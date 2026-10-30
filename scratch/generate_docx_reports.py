import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, bold=False, italic=False, color=None, font_size=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if font_size:
        run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = color
    return p

def format_inline_runs(paragraph, text, base_font_size=11, base_color=None):
    # Splits text by bold (**...**) and inline code (`...`) and italic (*...*)
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            content = token[2:-2]
            run = paragraph.add_run(content)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            if base_color:
                run.font.color.rgb = base_color
        elif token.startswith('`') and token.endswith('`'):
            content = token[1:-1]
            run = paragraph.add_run(content)
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('*') and token.endswith('*') and not token.startswith('**'):
            content = token[1:-1]
            run = paragraph.add_run(content)
            run.italic = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            if base_color:
                run.font.color.rgb = base_color
        else:
            run = paragraph.add_run(token)
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            if base_color:
                run.font.color.rgb = base_color

def convert_md_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    # Page setup - Margins 1 inch (approx 2.54cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_block_lang = ''
    code_lines = []
    
    in_table = False
    table_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check code block
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_block_lang = line[3:].strip()
                code_lines = []
            else:
                in_code_block = False
                # Render code block
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F4F6F9")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                # Title bar of code block
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(2)
                hdr_run = cp.add_run(f"[{code_block_lang.upper() if code_block_lang else 'CODE'}]" + "\n")
                hdr_run.font.name = 'Calibri'
                hdr_run.font.size = Pt(9)
                hdr_run.font.bold = True
                hdr_run.font.color.rgb = RGBColor(100, 110, 120)
                
                full_code = "\n".join(code_lines)
                c_run = cp.add_run(full_code)
                c_run.font.name = 'Consolas'
                c_run.font.size = Pt(9.5)
                c_run.font.color.rgb = RGBColor(35, 45, 55)
                
                # Spacer
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(4)
                
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Check Table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                # Process collected table lines
                parse_and_insert_table(doc, table_lines)
                table_lines = []
                
        # Empty line
        if not line.strip():
            i += 1
            continue
            
        # Divider ---
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # Add horizontal line border
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="D0D7DE"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue
            
        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            format_inline_runs(p, line[2:].strip(), base_font_size=20, base_color=RGBColor(26, 54, 93))
            p.runs[0].bold = True
            i += 1
            continue
            
        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(5)
            format_inline_runs(p, line[3:].strip(), base_font_size=15, base_color=RGBColor(30, 70, 120))
            p.runs[0].bold = True
            i += 1
            continue
            
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            format_inline_runs(p, line[4:].strip(), base_font_size=12.5, base_color=RGBColor(40, 80, 140))
            p.runs[0].bold = True
            i += 1
            continue
            
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            format_inline_runs(p, line[5:].strip(), base_font_size=11.5, base_color=RGBColor(50, 90, 150))
            p.runs[0].bold = True
            i += 1
            continue
            
        # Quote Callout > ...
        if line.startswith('> '):
            quote_text = line[2:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = tbl.cell(0, 0)
            set_cell_background(c, "F8FAFC")
            set_cell_margins(c, top=80, bottom=80, left=160, right=140)
            tcPr = c._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>
                    <w:top w:val="none"/>
                    <w:right w:val="none"/>
                    <w:bottom w:val="none"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)
            qp = c.paragraphs[0]
            qp.paragraph_format.space_after = Pt(2)
            format_inline_runs(qp, quote_text, base_font_size=10.5, base_color=RGBColor(51, 65, 85))
            i += 1
            continue
            
        # Bullet list - or *
        if re.match(r'^(\-|\*)\s+', line):
            text_part = re.sub(r'^(\-|\*)\s+', '', line).strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, text_part, base_font_size=11)
            i += 1
            continue
            
        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)$', line)
        if num_match:
            num = num_match.group(1)
            text_part = num_match.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, text_part, base_font_size=11)
            i += 1
            continue
            
        # Normal Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        format_inline_runs(p, line.strip(), base_font_size=11)
        i += 1

    if in_table:
        parse_and_insert_table(doc, table_lines)
        
    doc.save(docx_path)
    print(f"Saved: {docx_path}")

def parse_and_insert_table(doc, lines):
    if len(lines) < 2:
        return
        
    # Split cells
    raw_rows = []
    for l in lines:
        cleaned = l.strip()
        if cleaned.startswith('|'):
            cleaned = cleaned[1:]
        if cleaned.endswith('|'):
            cleaned = cleaned[:-1]
        cols = [c.strip() for c in cleaned.split('|')]
        raw_rows.append(cols)
        
    if len(raw_rows) < 2:
        return
        
    header_row = raw_rows[0]
    # Check if row 1 is separator (e.g. :---:|:---|...)
    is_sep = all(set(c).issubset({'-', ':', ' '}) for c in raw_rows[1])
    data_rows = raw_rows[2:] if is_sep else raw_rows[1:]
    
    num_cols = len(header_row)
    table = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    
    # Style Header Row
    hdr_cells = table.rows[0].cells
    for c_idx in range(num_cols):
        cell = hdr_cells[c_idx]
        set_cell_background(cell, "1E3A8A") # Navy Blue
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        txt = header_row[c_idx] if c_idx < len(header_row) else ''
        format_inline_runs(p, txt, base_font_size=10, base_color=RGBColor(255, 255, 255))
        p.runs[0].bold = True if p.runs else False
        
    # Style Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx in range(num_cols):
            cell = row_cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            txt = row_data[c_idx] if c_idx < len(row_data) else ''
            # Handle HTML <br> tags in markdown cells
            txt_lines = txt.split('<br>')
            for sub_idx, sub_txt in enumerate(txt_lines):
                if sub_idx > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(2)
                format_inline_runs(p, sub_txt.strip(), base_font_size=9.5, base_color=RGBColor(30, 41, 59))
                
    # Apply table borders
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)
    
    # Add spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)

if __name__ == '__main__':
    base_dir = r'c:\Users\Vi-Hung\my_flutter_pj\XD_HT_LMS\lms_trung_tam_ngoai_ngu\docs'
    tasks = [
        ('Task1_Logical_View.md', 'Task1_Logical_View.docx'),
        ('Task2_Process_View.md', 'Task2_Process_View.docx'),
        ('Task3_Implementation_View.md', 'Task3_Implementation_View.docx'),
        ('Task4_Deployment_View.md', 'Task4_Deployment_View.docx')
    ]
    for md_name, docx_name in tasks:
        md_file = os.path.join(base_dir, md_name)
        docx_file = os.path.join(base_dir, docx_name)
        if os.path.exists(md_file):
            convert_md_to_docx(md_file, docx_file)
        else:
            print(f"Not found: {md_file}")
